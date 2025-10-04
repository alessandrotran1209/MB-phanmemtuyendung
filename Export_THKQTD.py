import pandas as pd
import os
import base64
import requests
import time
import openpyxl

from openpyxl.utils.dataframe import dataframe_to_rows



while True:
    try:
                a = os.listdir("C:\\MB\\THKQTD")
                print(a)
                env_delimiter = "YYYYY"  
                delimiter = "XXXXXXXXXX"   
                envr, rest_email = a[0].split(env_delimiter, 1)
                if envr == "DEV":
                        THKQTD_Import = str("https://prod-26.southeastasia.logic.azure.com:443/workflows/31cec5643c42487d8b0dd838a6fdd02c/triggers/manual/paths/invoke?api-version=2016-06-01&sp=%2Ftriggers%2Fmanual%2Frun&sv=1.0&sig=-b6HtxU0jBJUZ8VZcKJrjK2j44pnW6vg3Y2knGq3dA8")
                elif envr == "UAT":
                        THKQTD_Import = str("https://prod-38.southeastasia.logic.azure.com:443/workflows/3bc0defb0fa14742b4ba8da7fdd2fa63/triggers/manual/paths/invoke?api-version=2016-06-01&sp=%2Ftriggers%2Fmanual%2Frun&sv=1.0&sig=ZUhcBlNMQvtZIvbgt7ZASn7sPKJgwkbxyJw7XZm9f1A")
                elif envr == "PROD":
                        THKQTD_Import = str("https://prod-17.southeastasia.logic.azure.com:443/workflows/c3e6213471f54c3492feb8056da49907/triggers/manual/paths/invoke?api-version=2016-06-01&sp=%2Ftriggers%2Fmanual%2Frun&sv=1.0&sig=Mbv4qDVYyNkf0lGRLgXsBf9OyVDurUpb7KoTnJG57rc")
                owner, rest = rest_email.split(delimiter, 1)
                os.rename(r"C:\\MB\\THKQTD\\"+a[0],r"C:\\MB\\THKQTD\\"+rest)
                a[0] = rest[:-4]
                print(a[0])
                
                file = f"C:\\MB\\THKQTD_Template\\SME.xlsm" ## Your template file
                wb = openpyxl.load_workbook(filename=file, read_only=False, keep_vba=True)
                ws = wb.active
                
                
                print('1')
                df = pd.read_csv(f"C:\\MB\\THKQTD\\{a[0]}.csv",encoding="utf8")
                
                rows = dataframe_to_rows(df, index=False, header=None)
                for r_idx,row in enumerate(rows, 16):
                        for c_idx, value in enumerate(row, 1):
                                ws.cell(row=r_idx, column=c_idx, value=value) ##Add the data
                xxx = 16+len(df)
                ws.delete_rows(16+len(df),15-len(df))
                ws.merge_cells(f"A{xxx}:P{xxx}")
                ws.merge_cells(f"I{xxx+1}:P{xxx+1}")
                ws.merge_cells(f"I{xxx+7}:P{xxx+7}")
                for i in range(0,8):
                        ws.row_dimensions[xxx+i].height = 20
                ws.print_area = f"A1:P{23+len(df)}"
                                

                wb.save(f"C:\\MB\\THKQTD\\{a[0]}.xlsm") ##Your output file
                print('dfdf')

                wb.close()
                with open(f"C:\\MB\\THKQTD\\{a[0]}.xlsm", "rb") as xlsxbinary1:
                    xlsxfile = base64.b64encode(xlsxbinary1.read()).decode('ascii')
                datasms = {"owner":str(owner),"file":xlsxfile}
                print("Sending to cloud")
                requests.post(THKQTD_Import, json = datasms)
                os.remove(f"C:\\MB\\THKQTD\\{a[0]}.xlsm")
                os.remove(f"C:\\MB\\THKQTD\\{a[0]}.csv")

                print("Done")
    except Exception as error:
        print(error)
        print(str(time.ctime()))
        time.sleep(1)
        pass

            


